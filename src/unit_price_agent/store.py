"""Excel 與照片資料夾的讀寫。

防護 4（寫進 Excel 的文字一律當純文字）在這個檔案。
紀錄與照片只由佇列的 worker 存取（見 bot.py），所以這裡不處理同時讀寫。
"""

import json
import os
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.worksheet.worksheet import Worksheet

from .validate import ValidationError, build_photo_filename, ensure_inside

HEADERS = ["紀錄編號", "日期時間", "上傳者 Discord ID", "上傳者", "類別", "品名", "品牌", "價格",
           "容量", "單位", "換算單價", "店家或地點", "照片位置", "回覆訊息 ID"]
(COL_ID, COL_WHEN, COL_USER_ID, COL_USER, COL_CATEGORY, COL_NAME, COL_BRAND, COL_PRICE,
 COL_AMOUNT, COL_UNIT, COL_UNIT_PRICE, COL_STORE, COL_PHOTO, COL_REPLY_ID) = range(1, 15)
TIME_FORMAT = "%Y-%m-%d %H:%M:%S"


class StoreError(Exception):
    """讀寫紀錄或照片失敗（例如 Excel 被打開鎖住）。訊息不含使用者給的內容。"""


@dataclass(frozen=True)
class Record:
    record_id: int
    when: str
    user_id: str       # Discord ID 存成文字：數字太長，存成數字 Excel 會丟精度
    user_name: str
    category: str
    name: str
    brand: str
    price: float
    amount: float
    unit: str
    unit_price: float
    store: str
    photo: str         # 相對於 DATA_DIR 的路徑
    reply_id: str


# ==== 防護 4：Excel 文字欄位一律存成純文字　開始 ====
def set_text_cell(ws: Worksheet, row: int, column: int, text: str) -> Cell:
    """把文字寫進儲存格，並明確設成「純文字」。

    Excel 會把以 = + - @ 開頭（或開頭是 Tab、換行）的文字當成公式。這裡不靠補撇號，
    而是把儲存格的資料型別設成文字（'s'）、格式設成文字（'@'），內容原樣保留。
    所有文字欄位（品名、品牌、店家、類別、上傳者顯示名稱、照片位置）都要走這個函式。
    """
    text = ILLEGAL_CHARACTERS_RE.sub("", str(text))  # openpyxl 不收的控制字元
    cell = ws.cell(row=row, column=column)
    cell.value = text  # openpyxl 這時可能把 '=' 開頭的字串判成公式
    cell.data_type = "s"  # 明確改回文字
    cell.number_format = "@"
    return cell
# ==== 防護 4 結束 ====


def load_state(path: Path) -> dict:
    """讀 state.json（限流計數、最後一次發週報的時間）。不存在或壞掉就當空的。"""
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return data if isinstance(data, dict) else {}


def save_state(path: Path, state: dict) -> None:
    """先寫暫存檔再換名，避免寫到一半被中斷留下壞檔。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(state, ensure_ascii=False), encoding="utf-8")
    os.replace(tmp, path)


def _text(value: object) -> str:
    return "" if value is None else str(value)


def _record_from_row(row: tuple) -> Record:
    return Record(
        record_id=int(row[COL_ID - 1]), when=_text(row[COL_WHEN - 1]),
        user_id=_text(row[COL_USER_ID - 1]), user_name=_text(row[COL_USER - 1]),
        category=_text(row[COL_CATEGORY - 1]), name=_text(row[COL_NAME - 1]),
        brand=_text(row[COL_BRAND - 1]), price=float(row[COL_PRICE - 1]),
        amount=float(row[COL_AMOUNT - 1]), unit=_text(row[COL_UNIT - 1]),
        unit_price=float(row[COL_UNIT_PRICE - 1]), store=_text(row[COL_STORE - 1]),
        photo=_text(row[COL_PHOTO - 1]), reply_id=_text(row[COL_REPLY_ID - 1]))


class RecordStore:
    """DATA_DIR 底下的 records.xlsx 和 photos/。"""

    def __init__(self, data_dir: Path):
        self.data_dir = data_dir
        self.records_path = data_dir / "records.xlsx"
        self.photos_dir = data_dir / "photos"

    # ---- 讀 ----
    def all_records(self) -> list[Record]:
        if not self.records_path.exists():
            return []
        try:
            wb = load_workbook(self.records_path, read_only=True)
            try:
                rows = list(wb.active.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True))
            finally:
                wb.close()
            return [_record_from_row(r) for r in rows if r[COL_ID - 1] is not None]
        except (OSError, ValueError, TypeError, KeyError) as exc:
            raise StoreError("讀不到 records.xlsx") from exc

    def user_records(self, user_id: int | str) -> list[Record]:
        return [r for r in self.all_records() if r.user_id == str(user_id)]

    def user_categories(self, user_id: int | str) -> list[str]:
        """這個人用過的類別名稱（依第一次出現的順序，不重複）。"""
        seen: dict[str, None] = {}
        for r in self.user_records(user_id):
            seen.setdefault(r.category)
        return list(seen)

    # ---- 寫 ----
    def _write_photo(self, when: datetime, name: str, ext: str, data: bytes) -> Path:
        name_ = build_photo_filename(when, name, ext)
        self.photos_dir.mkdir(parents=True, exist_ok=True)
        target = ensure_inside(self.photos_dir, self.photos_dir / name_)
        n = 2
        while target.exists():  # 同一秒、同品名：加編號，不覆蓋舊的
            target = ensure_inside(self.photos_dir, self.photos_dir / f"{Path(name_).stem}-{n}{ext}")
            n += 1
        target.write_bytes(data)
        return target

    def _save_workbook(self, wb: Workbook) -> None:
        """先存成暫存檔再換名。Excel 開著鎖住 records.xlsx 時換名會失敗，原檔不動。"""
        tmp = self.records_path.with_suffix(".xlsx.tmp")
        try:
            wb.save(tmp)
            os.replace(tmp, self.records_path)
        finally:
            tmp.unlink(missing_ok=True)

    def _open_workbook(self) -> Workbook:
        if self.records_path.exists():
            return load_workbook(self.records_path)
        wb = Workbook()
        wb.active.append(HEADERS)
        return wb

    def add_record(self, *, when: datetime, user_id: int, user_name: str, category: str, name: str,
                   brand: str, price: float, amount: float, unit: str, unit_price: float,
                   store: str, photo_bytes: bytes, photo_ext: str, reply_id: int | str = "") -> Record:
        """存照片、加一列。任何一步失敗都不留照片也不留列，丟 StoreError。"""
        photo_path = None
        try:
            self.data_dir.mkdir(parents=True, exist_ok=True)
            photo_path = self._write_photo(when, name, photo_ext, photo_bytes)
            wb = self._open_workbook()
            ws = wb.active
            record_id = max([ws.cell(row=r, column=COL_ID).value or 0
                             for r in range(2, ws.max_row + 1)], default=0) + 1
            row = ws.max_row + 1
            relative_photo = photo_path.relative_to(self.data_dir.resolve()).as_posix()
            ws.cell(row=row, column=COL_ID, value=record_id)
            set_text_cell(ws, row, COL_WHEN, when.strftime(TIME_FORMAT))
            set_text_cell(ws, row, COL_USER_ID, str(user_id))
            set_text_cell(ws, row, COL_USER, user_name)
            set_text_cell(ws, row, COL_CATEGORY, category)
            set_text_cell(ws, row, COL_NAME, name)
            set_text_cell(ws, row, COL_BRAND, brand)
            ws.cell(row=row, column=COL_PRICE, value=price)
            ws.cell(row=row, column=COL_AMOUNT, value=amount)
            set_text_cell(ws, row, COL_UNIT, unit)
            ws.cell(row=row, column=COL_UNIT_PRICE, value=unit_price)
            set_text_cell(ws, row, COL_STORE, store)
            set_text_cell(ws, row, COL_PHOTO, relative_photo)
            set_text_cell(ws, row, COL_REPLY_ID, str(reply_id))
            self._save_workbook(wb)
        except (OSError, ValueError, ValidationError, KeyError) as exc:
            if photo_path is not None:
                photo_path.unlink(missing_ok=True)
            raise StoreError("寫入紀錄失敗") from exc
        return Record(record_id, when.strftime(TIME_FORMAT), str(user_id), user_name, category, name,
                      brand, price, amount, unit, unit_price, store, relative_photo, str(reply_id))

    def set_reply_id(self, record_id: int, reply_id: int | str) -> None:
        """回覆送出後，把回覆訊息 ID 記進那一列（修正時靠它找到是哪一筆）。"""
        try:
            wb = self._open_workbook()
            ws = wb.active
            for r in range(2, ws.max_row + 1):
                if ws.cell(row=r, column=COL_ID).value == record_id:
                    set_text_cell(ws, r, COL_REPLY_ID, str(reply_id))
                    self._save_workbook(wb)
                    return
        except (OSError, ValueError, KeyError) as exc:
            raise StoreError("更新回覆訊息 ID 失敗") from exc
        raise StoreError("找不到這筆紀錄")
