import pytest
from openpyxl import Workbook, load_workbook

from unit_price_agent.store import set_text_cell

DANGEROUS = ["=1+1", "+1", "-1", "@SUM(A1)", "\t=1+1", "\n=1+1", "\r=1", "=HYPERLINK(\"http://x\")"]


@pytest.mark.parametrize("text", DANGEROUS + ["普通文字"])
def test_text_cells_stay_plain_text_after_save_and_reload(tmp_path, text):
    wb = Workbook()
    set_text_cell(wb.active, 1, 1, text)
    path = tmp_path / "t.xlsx"
    wb.save(path)
    cell = load_workbook(path).active.cell(row=1, column=1)
    assert cell.data_type == "s"
    # XML 會把 \r 讀成 \n，這是檔案格式的特性，不影響它還是純文字
    assert cell.value == text.replace("\r", "\n")


def test_text_cell_uses_text_number_format():
    cell = set_text_cell(Workbook().active, 1, 1, "=1+1")
    assert cell.number_format == "@" and cell.data_type == "s"


def test_illegal_control_characters_do_not_crash():
    cell = set_text_cell(Workbook().active, 1, 1, "品名\x00\x07")
    assert cell.value == "品名"


# ---- RecordStore ----
from datetime import datetime  # noqa: E402

from unit_price_agent.store import HEADERS, RecordStore, StoreError  # noqa: E402

PNG = b"\x89PNG\r\n\x1a\n" + b"0" * 16
WHEN = datetime(2026, 1, 2, 3, 4, 5)


def add(store, name="鮮乳", user_id=123456789012345678, **kw):
    args = dict(when=WHEN, user_id=user_id, user_name="小明", category="鮮奶", name=name, brand="光泉",
                price=89, amount=936, unit="ml", unit_price=9.5, store="", photo_bytes=PNG, photo_ext=".png")
    args.update(kw)
    return store.add_record(**args)


def test_add_record_writes_14_columns_photo_and_reads_back(tmp_path):
    store = RecordStore(tmp_path)
    rec = add(store)
    ws = load_workbook(tmp_path / "records.xlsx").active
    assert [c.value for c in ws[1]] == HEADERS and len(HEADERS) == 14
    assert (tmp_path / rec.photo).read_bytes() == PNG
    assert rec.photo == "photos/20260102-030405_鮮乳.png"
    [back] = store.all_records()
    assert back == rec and back.user_id == "123456789012345678"  # 長 ID 沒有被丟精度


def test_record_ids_increase_and_users_are_separate(tmp_path):
    store = RecordStore(tmp_path)
    assert [add(store).record_id, add(store, user_id=2).record_id] == [1, 2]
    assert len(store.user_records(2)) == 1 and store.user_categories(3) == []


def test_same_name_same_second_does_not_overwrite_photo(tmp_path):
    store = RecordStore(tmp_path)
    assert add(store).photo != add(store).photo


def test_hostile_text_is_stored_as_plain_text_and_photo_stays_inside(tmp_path):
    store = RecordStore(tmp_path)
    rec = add(store, name="../../evil", brand="=1+1")
    assert ".." not in rec.photo and not (tmp_path.parent / "evil.png").exists()
    cell = load_workbook(tmp_path / "records.xlsx").active.cell(row=2, column=7)
    assert cell.data_type == "s" and cell.value == "=1+1"


def test_set_reply_id(tmp_path):
    store = RecordStore(tmp_path)
    rec = add(store)
    store.set_reply_id(rec.record_id, 999)
    assert store.all_records()[0].reply_id == "999"


def test_failed_write_leaves_no_photo_and_no_row(tmp_path, monkeypatch):
    store = RecordStore(tmp_path)
    add(store)

    def locked(*a, **k):
        raise PermissionError("locked")
    monkeypatch.setattr("unit_price_agent.store.os.replace", locked)
    with pytest.raises(StoreError):
        add(store, name="另一個")
    assert len(store.all_records()) == 1
    assert len(list((tmp_path / "photos").iterdir())) == 1
    assert not list(tmp_path.glob("*.tmp"))


def test_find_by_reply_id_and_update_record(tmp_path):
    store = RecordStore(tmp_path)
    rec = add(store, reply_id=555)
    add(store, user_id=2, reply_id=556)
    assert store.find_by_reply_id(555) == rec and store.find_by_reply_id(1) is None
    updated = store.update_record(rec.record_id, name="=1+1", price=99.0, unit_price=1.5)
    assert (updated.name, updated.price, updated.unit_price, updated.photo) == ("=1+1", 99.0, 1.5, rec.photo)
    assert store.find_by_reply_id(556).price == 89  # 別筆不動
    cell = load_workbook(tmp_path / "records.xlsx").active.cell(row=2, column=6)
    assert cell.data_type == "s" and cell.value == "=1+1"


def test_update_record_rejects_unknown_fields_and_missing_ids(tmp_path):
    store = RecordStore(tmp_path)
    add(store)
    with pytest.raises(StoreError):
        store.update_record(1, photo="x")
    with pytest.raises(StoreError):
        store.update_record(1, reply_id="1")
    with pytest.raises(StoreError):
        store.update_record(99, name="x")
    assert store.all_records()[0].record_id == 1
